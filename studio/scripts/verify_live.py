#!/usr/bin/env python3
"""Generate real, bounded Office deliverables through the deployed BMA app."""
import argparse
import io
import json
import time
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import quote

import httpx
from docx import Document
from openpyxl import Workbook
from pptx import Presentation


def verify(case):
    root = Path(__file__).resolve().parents[2]
    deployment = json.loads((root / "artifacts/studio/deployment.json").read_text())
    token = json.loads((root / deployment["secret_file"]).read_text())["app_access_token"]
    output = root / "artifacts/studio/live"
    output.mkdir(parents=True, exist_ok=True)
    base, headers = deployment["outputs"]["WebUrl"], {"Authorization": "Bearer " + token}
    request_id = str(uuid.uuid4())
    request = {"request_id": request_id, "language": "ko", "style": "business",
               "outputs": ["chart", "pptx", "docx"]}
    with httpx.Client(timeout=45) as client:
        config = client.get(base + "/api/config", headers=headers)
        config.raise_for_status()
        assert len(config.json()["styles"]) == 3
        if case == "excel":
            wb = Workbook()
            ws = wb.active
            ws.title = "월별 실적"
            ws.append(["월", "매출"])
            for row in [("1월", 120), ("2월", 150), ("3월", 180), ("4월", 160), ("5월", 220), ("6월", 250)]:
                ws.append(row)
            data = io.BytesIO()
            wb.save(data)
            response = client.post(base + "/api/uploads", content=data.getvalue(),
                                   headers={**headers, "Content-Type": "application/octet-stream",
                                            "X-File-Name": quote("월별_매출.xlsx")})
            response.raise_for_status()
            request["upload_id"] = response.json()["upload"]["id"]
            request["prompt"] = "첨부한 월별 매출로 막대 차트 한 개를 만들고, 차트를 포함한 간단한 3장 발표자료와 Word 보고서를 만들어 주세요. 샘플 데이터임을 표시하고 제공되지 않은 원인이나 전망은 쓰지 마세요."
        elif case == "presentation":
            first = json.loads((output / "excel-summary.json").read_text())
            request.update(source_job_id=first["job_id"], style="presentation",
                           prompt="같은 월별 매출 자료를 발표용 스타일로 다시 만들어 주세요. 실제 값의 막대 차트 한 개와 짧은 3장 PPT, Word 보고서가 필요합니다.")
        else:
            request.update(style="aws-desk", outputs=["diagram", "pptx", "docx"],
                           notes="신입사원 온보딩 업무 메모\n1. 인사팀이 입사 서류를 접수한다.\n2. IT팀이 계정과 장비를 준비한다.\n3. 팀장이 업무 소개와 첫 주 일정을 안내한다.\n4. 멘토가 첫 주 질의응답을 진행한다.\n5. 인사팀이 적응 상태를 확인한다.\n이 메모는 데모용 예시다.",
                           prompt="이 메모를 업무 흐름도로 시각화하고, 흐름도를 포함한 짧은 PPT와 Word 안내서를 만들어 주세요. 메모에 없는 일정이나 담당자를 추가하지 마세요.")
        created = client.post(base + "/api/jobs", json=request, headers=headers)
        created.raise_for_status()
        assert created.json()["job"]["id"] == request_id
    start = time.monotonic()
    events, final = [], None
    with httpx.Client(timeout=httpx.Timeout(45, connect=10)) as client:
        with client.stream("GET", base + f"/api/jobs/{request_id}/events", headers=headers) as response:
            response.raise_for_status()
            for line in response.iter_lines():
                if time.monotonic() - start > 350:
                    raise TimeoutError("Studio validation deadline")
                if not line.startswith("data: "):
                    continue
                event = json.loads(line[6:])
                events.append(event)
                if event["type"] == "snapshot":
                    job = event["job"]
                    print(json.dumps({"case": case, "status": job["status"], "phase": job["phase"],
                                      "files": len(job["artifacts"])}), flush=True)
                if event["type"] == "complete":
                    final = event["job"]
        (output / f"{case}-{request_id}.json").write_text(
            json.dumps({"request": request, "events": events, "result": final}, ensure_ascii=False, indent=2) + "\n")
        assert final and final["status"] == "complete", final
        assert not final["errors"]
        formats, chart_values = set(), []
        for item in final["artifacts"]:
            artifact = client.get(base + f"/api/jobs/{request_id}/artifacts/{item['id']}", headers=headers)
            artifact.raise_for_status()
            assert len(artifact.content) == item["size"]
            path = output / f"{case}-{item['filename']}"
            path.write_bytes(artifact.content)
            if item["filename"].endswith(".pptx"):
                formats.add("pptx")
                with zipfile.ZipFile(io.BytesIO(artifact.content)) as z:
                    assert z.testzip() is None
                ppt = Presentation(io.BytesIO(artifact.content))
                assert 1 <= len(ppt.slides) <= 12
                for slide in ppt.slides:
                    for shape in slide.shapes:
                        if shape.has_chart:
                            chart_values.extend(shape.chart.series[0].values)
            elif item["filename"].endswith(".docx"):
                formats.add("docx")
                doc = Document(io.BytesIO(artifact.content))
                assert any(p.text.strip() for p in doc.paragraphs)
                assert doc.inline_shapes
            elif item["mime"] == "text/html":
                formats.add("html")
                assert "<script" not in artifact.text.lower()
            elif item["kind"] in {"chart", "diagram"}:
                formats.add(item["kind"])
        required = {"pptx", "docx", "html", "diagram" if case == "notes" else "chart"}
        assert required <= formats
        if case != "notes":
            assert sorted(chart_values) == [120, 150, 160, 180, 220, 250], chart_values
    proof = {"status": "passed", "case": case, "job_id": request_id, "model": deployment["model"],
             "style": request["style"], "formats": sorted(formats), "artifacts": len(final["artifacts"]),
             "chart_values": chart_values, "elapsed_seconds": round(time.monotonic() - start, 2),
             "verified_at": datetime.now(timezone.utc).isoformat(), "warnings": final["warnings"]}
    (output / f"{case}-summary.json").write_text(json.dumps(proof, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(proof, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--case", choices=["excel", "presentation", "notes"], default="excel")
    verify(parser.parse_args().case)
