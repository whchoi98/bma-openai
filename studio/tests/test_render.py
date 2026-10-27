"""Behavior and file-format contracts for local document rendering."""
import copy
import importlib
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from PIL import Image


class PreviewParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags = []
        self.attrs = []

    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)
        self.attrs.extend(attrs)


CHART = {
    "title": "분기별 매출 <검토>", "chart_type": "bar",
    "categories": ["서울", "부산", "인천"], "values": [10, None, 30],
    "x_label": "지역", "y_label": "매출", "source": {"file": "매출.csv", "sheet": "Sheet1"},
}
DIAGRAM = {
    "title": "자료 처리 흐름", "layout": "flow",
    "nodes": [{"id": "a", "label": "자료 <수집>"}, {"id": "b", "label": "검토 & 승인"}],
    "edges": [{"source": "a", "target": "b", "label": "전달"}],
}
BLUEPRINT = {
    "title": "사업 결과 <script>alert(1)</script>", "subtitle": "2026년 현황",
    "sections": [{"heading": "요약", "paragraphs": ["자료를 분석했습니다. <img src=x onerror=alert(1)>"],
                  "bullets": ["원본 수치 보존", "검토 후 공유"], "asset_ids": ["chart_1"]},
                 {"heading": "처리 과정", "asset_ids": ["diagram_1"]}],
}


class RenderTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            cls.render = importlib.import_module("studio.document_tools.render")
        except ModuleNotFoundError:
            cls.render = None

    def setUp(self):
        self.assertIsNotNone(self.render, "The document renderer must be implemented")
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.out = Path(self.temp.name)

    def file(self, asset, mime):
        return Path(next(f["path"] for f in asset["files"] if f["mime"] == mime))

    def test_chart_png_svg_unicode_and_missing_values(self):
        for kind in ("bar", "line", "pie", "scatter"):
            with self.subTest(kind=kind):
                spec = copy.deepcopy(CHART)
                spec["chart_type"] = kind
                if kind == "scatter":
                    spec["values"] = [10, 20, 30]
                    spec["x_values"] = [1, 2, 4]
                asset = self.render.render_chart(spec, self.out, asset_id=kind)
                self.assertEqual(asset["spec"]["values"], spec["values"])
                with Image.open(self.file(asset, "image/png")) as image:
                    self.assertGreaterEqual(image.width, 1000)
                    self.assertGreaterEqual(image.height, 500)
                    self.assertGreater(len(image.convert("RGB").getcolors(1_000_000)), 20)
                svg = self.file(asset, "image/svg+xml").read_text()
                root = ET.fromstring(svg)
                self.assertEqual(root.tag, "{http://www.w3.org/2000/svg}svg")
                self.assertNotIn("<script", svg)
                self.assertIn("분기별 매출", svg)
                if kind != "scatter":
                    self.assertIn("서울", svg)

    def test_parser_chart_metadata_survives_rendering(self):
        from studio.document_tools.input_data import aggregate_chart, inspect_source
        source = inspect_source("지역,매출\n서울,10\n부산,20\n".encode(), "매출.csv")
        spec = aggregate_chart(source, "sheet_0", "c0", "c1")
        asset = self.render.render_chart(spec, self.out)
        self.assertEqual(asset["spec"]["source"]["x_column"], {"id": "c0", "name": "지역", "type": "string"})
        self.assertEqual(asset["spec"]["values"], [10.0, 20.0])

    def test_long_excel_labels_render_with_disclosure_and_original_workbook_data(self):
        import io

        import openpyxl

        from studio.document_tools.input_data import aggregate_chart, inspect_source
        x_header = "지역별 실적 분류 " * 24
        y_header = "분기별 매출 원본 " * 24
        first = "같은 접두사의 한글 범주 " * 24 + "첫째"
        second = "같은 접두사의 한글 범주 " * 24 + "둘째"
        book = openpyxl.Workbook()
        book.active.append([x_header, y_header])
        book.active.append([first, 15])
        book.active.append([second, 25])
        buffer = io.BytesIO()
        book.save(buffer)
        source = inspect_source(buffer.getvalue(), "긴원본.xlsx")
        spec = aggregate_chart(source, "sheet_0", "c0", "c1")
        asset = self.render.render_chart(spec, self.out)
        self.assertEqual(asset["spec"]["categories"], [first, second])
        self.assertEqual(asset["spec"]["x_label"], x_header)
        self.assertEqual(asset["spec"]["y_label"], y_header)
        self.assertEqual(asset["spec"]["source"]["x_column"]["name"], x_header)
        self.assertEqual(asset["spec"]["values"], [15, 25])
        display = asset["spec"]["display"]
        self.assertLess(len(display["x_label"]), 161)
        self.assertLess(len(display["y_label"]), 161)
        self.assertTrue(all(len(label) < 161 and "…" in label for label in display["categories"]))
        self.assertNotEqual(display["categories"][0], display["categories"][1])
        self.assertIn("categories[0]", display["shortened"])
        self.assertIn("x_label", display["shortened"])
        self.assertTrue(any("shortened" in warning for warning in asset["spec"]["warnings"]))
        self.assertLessEqual(len(asset["title"]), 180)
        svg = self.file(asset, "image/svg+xml").read_text()
        self.assertIn("Labels shortened", svg)
        ET.fromstring(svg)
        with Image.open(self.file(asset, "image/png")) as image:
            image.verify()
        bp = {"title": "긴 이름 보존", "sections": [{"heading": "차트", "asset_ids": ["chart_1"]}]}
        result = self.render.render_documents(bp, {"chart_1": asset}, self.out)
        with zipfile.ZipFile(self.file(result, "application/vnd.openxmlformats-officedocument.presentationml.presentation")) as zf:
            embedded = next(n for n in zf.namelist() if n.startswith("ppt/embeddings/"))
            workbook = openpyxl.load_workbook(io.BytesIO(zf.read(embedded)))
            originals = [row[3] for row in workbook["Source labels"].iter_rows(min_row=2, values_only=True)]
            for label in (first, second, x_header, y_header):
                self.assertIn(label, originals)
            self.assertEqual(workbook["Sheet1"]["B2"].value, 15)
            self.assertEqual(workbook["Sheet1"]["B3"].value, 25)
            workbook.close()
            slides = "".join(zf.read(n).decode() for n in zf.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
            self.assertIn("Labels shortened", slides)
        self.assertIn("Labels shortened", self.file(result, "text/html").read_text())

    def test_original_labels_over_excel_cell_limit_are_split_without_loss(self):
        import io

        import openpyxl
        original = "자료" * 20000
        spec = {**CHART, "categories": [original], "values": [12]}
        asset = self.render.render_chart(spec, self.out)
        bp = {"title": "긴 원본", "sections": [{"heading": "차트", "asset_ids": ["chart_1"]}]}
        result = self.render.render_documents(bp, {"chart_1": asset}, self.out, formats=("pptx",))
        with zipfile.ZipFile(self.file(result, "application/vnd.openxmlformats-officedocument.presentationml.presentation")) as zf:
            embedded = next(n for n in zf.namelist() if n.startswith("ppt/embeddings/"))
            workbook = openpyxl.load_workbook(io.BytesIO(zf.read(embedded)))
            rows = list(workbook["Source labels"].iter_rows(min_row=2, values_only=True))
            chunks = [row[3] for row in rows if row[0] == "category" and row[1] == 1]
            self.assertEqual("".join(chunks), original)
            self.assertTrue(all(len(chunk) <= 32767 for chunk in chunks))
            self.assertEqual(asset["spec"]["categories"], [original])
            workbook.close()

    def test_svg_is_self_contained_and_treats_markup_as_text(self):
        spec = {**CHART, "title": "$not-a-math-command$ <script>x</script>"}
        asset = self.render.render_chart(spec, self.out)
        svg = self.file(asset, "image/svg+xml").read_text()
        self.assertNotIn("<!DOCTYPE", svg)
        root = ET.fromstring(svg)
        for node in root.iter():
            self.assertNotEqual(node.tag.rsplit("}", 1)[-1], "script")
            for key, value in node.attrib.items():
                if key.rsplit("}", 1)[-1] == "href":
                    self.assertTrue(value.startswith("#"))

    def test_three_styles_affect_rendered_image(self):
        outputs = []
        for style in ("business", "aws-desk", "presentation"):
            asset = self.render.render_chart(CHART, self.out, style=style, asset_id=style)
            outputs.append(self.file(asset, "image/png").read_bytes())
        self.assertEqual(len(set(outputs)), 3)
        with self.assertRaises(ValueError):
            self.render.render_chart(CHART, self.out, style="unknown")

    def test_diagram_layouts_escape_labels(self):
        for layout in ("flow", "hierarchy", "timeline"):
            spec = copy.deepcopy(DIAGRAM)
            spec["layout"] = layout
            spec["nodes"][0]["label"] = '<script>alert("x")</script> & 입력'
            asset = self.render.render_diagram(spec, self.out, asset_id=layout)
            svg = self.file(asset, "image/svg+xml").read_text()
            ET.fromstring(svg)
            self.assertNotIn("<script", svg)
            self.assertIn("&lt;script&gt;", svg)
            with Image.open(self.file(asset, "image/png")) as image:
                image.verify()

    def test_office_native_charts_shapes_and_safe_preview(self):
        for style in ("business", "aws-desk", "presentation"):
            with self.subTest(style=style):
                folder = self.out / style
                chart = self.render.render_chart(CHART, folder, style)
                diagram = self.render.render_diagram(DIAGRAM, folder, style)
                result = self.render.render_documents(BLUEPRINT, {"chart_1": chart, "diagram_1": diagram}, folder, style,
                                                      source_notes=["출처: 매출.csv / 원본 3행"])
                pptx = self.file(result, "application/vnd.openxmlformats-officedocument.presentationml.presentation")
                with zipfile.ZipFile(pptx) as zf:
                    names = zf.namelist()
                    self.assertIn("ppt/charts/chart1.xml", names)
                    self.assertTrue(any(n.startswith("ppt/embeddings/") for n in names))
                    chart_xml = zf.read("ppt/charts/chart1.xml").decode()
                    self.assertIn("<c:barChart>", chart_xml)
                    self.assertIn("서울", chart_xml)
                    self.assertIn("<c:v>30</c:v>", chart_xml)
                    slides = [n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
                    self.assertEqual(len(slides), 3)
                    text = "".join(zf.read(n).decode() for n in slides)
                    self.assertIn("자료 &lt;수집&gt;", text)
                    self.assertIn("<p:cxnSp>", text)
                    self.assertIn("원본 3행", text)
                    self.assertIn("NanumSquare", text)
                    for name in names:
                        if name.endswith(".rels"):
                            self.assertNotIn('TargetMode="External"', zf.read(name).decode())
                docx = self.file(result, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                with zipfile.ZipFile(docx) as zf:
                    xml = zf.read("word/document.xml").decode()
                    self.assertIn("자료를 분석했습니다.", xml)
                    self.assertIn("원본 3행", xml)
                    self.assertGreaterEqual(len([n for n in zf.namelist() if n.startswith("word/media/")]), 2)
                html = self.file(result, "text/html").read_text()
                parser = PreviewParser()
                parser.feed(html)
                self.assertNotIn("script", parser.tags)
                self.assertFalse(any(k.lower().startswith("on") for k, _ in parser.attrs))
                self.assertTrue(all(v.startswith("data:image/png;base64,") for k, v in parser.attrs if k == "src"))
                self.assertIn("원본 3행", html)
                self.assertIn("content preview", html.lower())
                self.assertIn("&lt;script&gt;", html)

    def test_native_line_pie_and_scatter_chart_parts(self):
        for kind, tag in (("line", "lineChart"), ("pie", "pieChart"), ("scatter", "scatterChart")):
            spec = {**CHART, "chart_type": kind, "values": [10, 20, 30], "x_values": [1, 2, 3]}
            asset = self.render.render_chart(spec, self.out)
            bp = {"title": "Chart", "sections": [{"heading": kind, "asset_ids": ["chart_1"]}]}
            result = self.render.render_documents(bp, {"chart_1": asset}, self.out, formats=("pptx",))
            with zipfile.ZipFile(self.file(result, "application/vnd.openxmlformats-officedocument.presentationml.presentation")) as zf:
                self.assertIn(f"<c:{tag}", zf.read("ppt/charts/chart1.xml").decode())

    def test_embedded_chart_workbook_keeps_formula_like_labels_literal(self):
        import io

        import openpyxl
        spec = {**CHART, "categories": ['=HYPERLINK("https://example.invalid")', "normal", "last"]}
        asset = self.render.render_chart(spec, self.out)
        bp = {"title": "x", "sections": [{"heading": "Data", "asset_ids": ["chart_1"]}]}
        result = self.render.render_documents(bp, {"chart_1": asset}, self.out, formats=("pptx",))
        with zipfile.ZipFile(self.file(result, "application/vnd.openxmlformats-officedocument.presentationml.presentation")) as zf:
            name = next(n for n in zf.namelist() if n.startswith("ppt/embeddings/"))
            book = openpyxl.load_workbook(io.BytesIO(zf.read(name)))
            sheet = book.active
            self.assertEqual(sheet["A2"].data_type, "s")
            self.assertEqual(sheet["A2"].value, '=HYPERLINK("https://example.invalid")')
            self.assertEqual(sheet["B2"].value, 10)
            self.assertIsNone(sheet["B3"].value)
            self.assertEqual(sheet["B4"].value, 30)
            book.close()

    def test_rejects_invalid_chart_data_and_ids(self):
        invalid = [
            {**CHART, "values": [10, float("inf"), 20]}, {**CHART, "values": [1]},
            {**CHART, "values": [None, None, None]}, {**CHART, "values": [True, 1, 2]},
            {**CHART, "chart_type": "pie", "values": [-1, 2, 3]},
            {**CHART, "chart_type": "scatter"},
            {**CHART, "categories": list(map(str, range(51))), "values": [1] * 51},
        ]
        for spec in invalid:
            with self.subTest(spec=spec), self.assertRaises(ValueError):
                self.render.render_chart(spec, self.out)
        for asset_id in ("../escape", "a/b", "/tmp/escape", "x.svg\" onload=x"):
            with self.assertRaises(ValueError):
                self.render.render_chart(CHART, self.out, asset_id=asset_id)

    def test_rejects_invalid_nodes_edges_and_unbounded_documents(self):
        invalid = [
            {**DIAGRAM, "nodes": []},
            {**DIAGRAM, "nodes": [{"id": "a", "label": "a"}] * 2},
            {**DIAGRAM, "edges": [{"source": "a", "target": "missing"}]},
            {**DIAGRAM, "nodes": [{"id": str(i), "label": str(i)} for i in range(21)]},
            {**DIAGRAM, "edges": DIAGRAM["edges"] * 31},
        ]
        for spec in invalid:
            with self.assertRaises(ValueError):
                self.render.render_diagram(spec, self.out)
        blueprints = [
            {"title": "x", "sections": [{"heading": "x"}] * 12},
            {"title": "x", "sections": [{"heading": "x", "paragraphs": ["x"] * 100}]},
            {"title": "x", "sections": [{"heading": "x", "paragraphs": ["x" * 50000]}]},
            {"title": "x", "sections": [{"heading": "x", "asset_ids": ["missing"]}]},
        ]
        for blueprint in blueprints:
            with self.assertRaises(ValueError):
                self.render.render_documents(blueprint, {}, self.out)

    def test_rejects_content_that_cannot_fit_bounded_slide(self):
        asset = self.render.render_chart(CHART, self.out)
        bp = {"title": "x", "sections": [{"heading": "x", "paragraphs": ["가" * 600, "나" * 600], "asset_ids": ["chart_1"]}]}
        with self.assertRaises(ValueError):
            self.render.render_documents(bp, {"chart_1": asset}, self.out)

    def test_rejects_asset_path_escape_and_symlinks(self):
        asset = self.render.render_chart(CHART, self.out)
        bp = {"title": "x", "sections": [{"heading": "x", "asset_ids": ["chart_1"]}]}
        bad = copy.deepcopy(asset)
        bad["files"][0]["path"] = "/etc/passwd"
        with self.assertRaises(ValueError):
            self.render.render_documents(bp, {"chart_1": bad}, self.out)
        escaped = self.out / "escaped.png"
        escaped.symlink_to("/etc/passwd")
        bad["files"][0]["path"] = str(escaped)
        with self.assertRaises(ValueError):
            self.render.render_documents(bp, {"chart_1": bad}, self.out)
        target = self.out / "evil.png"
        target.symlink_to("/etc/passwd")
        with self.assertRaises(ValueError):
            self.render.render_chart(CHART, self.out, asset_id="evil")


if __name__ == "__main__":
    unittest.main()
