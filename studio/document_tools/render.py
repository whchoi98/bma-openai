"""Bounded local renderers. No network calls, model execution, or remote assets.

Numeric chart input is authoritative; missing observations remain missing. PPTX
uses native chart workbooks and editable diagram objects. HTML is a content
preview and intentionally does not promise desktop Office pagination.
"""
from __future__ import annotations

import base64
import html
import io
import math
import os
import re
import textwrap
import unicodedata
from functools import lru_cache
from pathlib import Path
from typing import Any

os.environ.setdefault("MPLCONFIGDIR", "/tmp/studio-matplotlib")
import matplotlib

matplotlib.use("Agg")
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as WordInches
from docx.shared import Pt as WordPt
from docx.shared import RGBColor as WordRGBColor
from matplotlib import font_manager
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement as PptxElement
from pptx.util import Inches, Pt
from xlsxwriter import Workbook

from .styles import DocumentStyle, get_style

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
# Input parser permits 50,000-character cells; aggregation may add a short
# header prefix (for example "Count of "). Display budgets are independent.
_MAX_SOURCE_LABEL = 50_032
_LABEL_WARNING = ("Long chart labels are shortened with … for display; original labels "
                  "are retained in the specification and the PowerPoint workbook's Source labels sheet.")
_LABEL_NOTICE = "Labels shortened (…); original labels are preserved."
_ID = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_-]{0,63}$")
_BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _text(value: Any, field: str, maximum: int, *, empty: bool = True) -> str:
    if not isinstance(value, str) or len(value) > maximum or _BAD_XML.search(value):
        raise ValueError(f"{field} must be text of at most {maximum} valid characters")
    if not empty and not value.strip():
        raise ValueError(f"{field} must not be empty")
    return value


def _identifier(value: Any) -> str:
    if not isinstance(value, str) or not _ID.fullmatch(value):
        raise ValueError("Asset ID must contain only letters, digits, underscore, or hyphen")
    return value


def _sequence(value: Any, field: str, maximum: int, *, minimum: int = 0) -> list:
    if not isinstance(value, (list, tuple)) or not minimum <= len(value) <= maximum:
        raise ValueError(f"{field} must contain {minimum} to {maximum} items")
    return list(value)


def _output_dir(outdir: Any) -> Path:
    root = Path(outdir).resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


def _target(root: Path, filename: str) -> Path:
    target = root / filename
    if target.is_symlink() or target.resolve().parent != root or (target.exists() and not target.is_file()):
        raise ValueError("Unsafe output path")
    return target


def _file(path: Path, mime: str) -> dict:
    return {"path": str(path), "filename": path.name, "mime": mime}


@lru_cache(maxsize=4)
def _font(path_override: str | None = None) -> str:
    candidates = [path_override, "/opt/studio/fonts/NanumSquareB.otf",
                  str(Path(__file__).resolve().parents[2] / "webapp/static/fonts/NanumSquareB.otf")]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            font_manager.fontManager.addfont(candidate)
            return font_manager.FontProperties(fname=candidate).get_name()
    return "DejaVu Sans"


def _rc(style: DocumentStyle) -> dict:
    return {"font.family": _font(os.environ.get("STUDIO_FONT_PATH")), "font.size": 12, "font.weight": "bold", "axes.labelweight": "bold",
            "text.color": style.ink, "axes.labelcolor": style.muted, "xtick.color": style.muted,
            "ytick.color": style.muted, "axes.edgecolor": style.grid, "axes.unicode_minus": False,
            "svg.fonttype": "path", "svg.hashsalt": "studio", "text.usetex": False, "text.parse_math": False}


def _number(value: Any, field: str, *, missing: bool = False):
    if value is None and missing:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{field} must contain finite numbers")
    try:
        if not math.isfinite(value) or abs(value) > 1e100:
            raise ValueError(f"{field} must contain finite numbers of magnitude at most 1e100")
    except OverflowError as exc:
        raise ValueError(f"{field} is too large") from exc
    return value


def _label_units(text: str) -> int:
    return sum(0 if unicodedata.combining(c) else 2 if unicodedata.east_asian_width(c) in "WF" else 1
               for c in text)


def _display_label(text: str, maximum: int, index: int | None = None) -> str:
    flat = " ".join(text.split())
    if _label_units(flat) <= maximum:
        return flat
    suffix = "…" + (f" [{index}]" if index is not None else "")
    budget, used, prefix = maximum - _label_units(suffix), 0, []
    for char in flat:
        width = _label_units(char)
        if used + width > budget:
            break
        prefix.append(char)
        used += width
    return "".join(prefix).rstrip() + suffix


def _wrap_label(text: str, maximum: int) -> str:
    lines, current, used = [], [], 0
    for char in text:
        width = _label_units(char)
        if current and used + width > maximum:
            lines.append("".join(current).strip())
            current, used = [], 0
        current.append(char)
        used += width
    if current:
        lines.append("".join(current).strip())
    return "\n".join(lines)


def _chart_display(spec: dict) -> dict:
    count = len(spec["categories"])
    category_budget = 48 if count <= 6 else 24 if count <= 20 else 12
    display = {"categories": [_display_label(label, category_budget, index + 1)
                              for index, label in enumerate(spec["categories"])],
               "x_label": _display_label(spec["x_label"], 64),
               "y_label": _display_label(spec["y_label"], 40),
               "title": _display_label(spec["title"], 96),
               "source_caption": _display_label(_source_caption(spec), 120),
               "shortened": []}
    for index, (original, shown) in enumerate(zip(spec["categories"], display["categories"])):
        if original != shown:
            display["shortened"].append(f"categories[{index}]")
    for field in ("x_label", "y_label", "title"):
        if spec[field] != display[field]:
            display["shortened"].append(field)
    if _source_caption(spec) != display["source_caption"]:
        display["shortened"].append("source_caption")
    return display


def _chart_spec(spec: Any) -> dict:
    if not isinstance(spec, dict):
        raise ValueError("Chart specification must be an object")
    kind = spec.get("chart_type", "bar")
    if not isinstance(kind, str) or kind not in {"bar", "line", "pie", "scatter"}:
        raise ValueError("Unsupported chart type")
    values = [_number(v, "values", missing=True) for v in _sequence(spec.get("values"), "values", 50, minimum=1)]
    categories = [_text(v, "category", _MAX_SOURCE_LABEL) for v in _sequence(spec.get("categories"), "categories", 50, minimum=1)]
    if len(values) != len(categories) or all(v is None for v in values):
        raise ValueError("Categories and numeric observations must match and contain data")
    if kind == "pie" and (any(v < 0 for v in values if v is not None) or not any(v and v > 0 for v in values)):
        raise ValueError("Pie charts require nonnegative data with a positive total")
    result = {"chart_type": kind, "categories": categories, "values": values,
              "x_label": _text(spec.get("x_label", ""), "x_label", _MAX_SOURCE_LABEL),
              "y_label": _text(spec.get("y_label", ""), "y_label", _MAX_SOURCE_LABEL)}
    result["title"] = _text(spec.get("title") or result["y_label"] or "Chart", "title", _MAX_SOURCE_LABEL, empty=False)
    if kind == "scatter":
        xs = [_number(v, "x_values") for v in _sequence(spec.get("x_values"), "x_values", 50, minimum=1)]
        if len(xs) != len(values):
            raise ValueError("Scatter x/y observations must match")
        result["x_values"] = xs
    # Preserve parser metadata only through bounded, typed fields.
    source = spec.get("source", {})
    if not isinstance(source, dict) or len(source) > 16:
        raise ValueError("Chart source metadata must be a bounded object")
    result["source"] = {}
    for key, value in source.items():
        _text(key, "source field", 80)
        if value is None or isinstance(value, (int, float, bool)):
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                _number(value, "source value")
            result["source"][key] = value
        elif isinstance(value, str):
            result["source"][key] = _text(value, "source value", 600)
        elif isinstance(value, dict):
            # Parser column schemas are shallow {id, name, type} objects.
            if len(value) > 8:
                raise ValueError("Chart column metadata is too large")
            result["source"][key] = {
                _text(k, "column metadata field", 80): _text(v, "column metadata value", _MAX_SOURCE_LABEL)
                for k, v in value.items()
            }
        elif isinstance(value, (list, tuple)):
            result["source"][key] = [_text(v, "source value", 160) for v in _sequence(value, "source values", 10)]
        else:
            raise ValueError("Unsupported chart source metadata")
    for name in ("missing_values", "invalid_values", "omitted_categories"):
        if name in spec:
            value = spec[name]
            if isinstance(value, bool) or not isinstance(value, int) or not 0 <= value <= 1_000_000:
                raise ValueError(f"Invalid {name}")
            result[name] = value
    result["warnings"] = [_text(w, "warning", 600) for w in _sequence(spec.get("warnings", []), "warnings", 40)]
    result["display"] = _chart_display(result)
    if result["display"]["shortened"] and _LABEL_WARNING not in result["warnings"]:
        result["warnings"].append(_LABEL_WARNING)
    return result


def _source_caption(spec: dict) -> str:
    source = spec.get("source", {})
    return " / ".join(str(source[key]) for key in ("file", "filename", "sheet") if source.get(key))


def _save_figure(fig: Figure, root: Path, asset_id: str, title: str) -> list[dict]:
    png = _target(root, f"{asset_id}.png")
    svg = _target(root, f"{asset_id}.svg")
    # Validate both destinations before the first write.
    fig.savefig(png, format="png", dpi=140, facecolor=fig.get_facecolor())
    fig.savefig(svg, format="svg", facecolor=fig.get_facecolor(), metadata={"Title": title, "Date": None})
    # Matplotlib emits a legacy external DTD. Remove it so the self-contained
    # SVG never offers an XML consumer a remote document to resolve.
    svg.write_text(re.sub(r"<!DOCTYPE[^>]*>", "", svg.read_text(encoding="utf-8")), encoding="utf-8")
    fig.clear()
    return [_file(png, "image/png"), _file(svg, "image/svg+xml")]


def render_chart(spec, outdir, style="business", asset_id="chart_1") -> dict:
    theme, spec, asset_id = get_style(style), _chart_spec(spec), _identifier(asset_id)
    root = _output_dir(outdir)
    display = spec["display"]
    with matplotlib.rc_context(_rc(theme)):
        fig = Figure(figsize=(12, 6.75), facecolor=theme.background)
        FigureCanvasAgg(fig)
        ax = fig.add_subplot(111)
        fig.subplots_adjust(left=.09, right=.95, top=.78, bottom=.24)
        ax.set_facecolor(theme.background)
        title = _wrap_label(display["title"], 48)
        fig.text(.065, .94, title, va="top", fontsize=23 if style == "presentation" else 20, weight="bold")
        fig.text(.065, .075, display["source_caption"], fontsize=9, color=theme.muted)
        if display["shortened"]:
            fig.text(.065, .03, _LABEL_NOTICE, fontsize=8, color=theme.muted)
        vals = [float(v) if v is not None else math.nan for v in spec["values"]]
        labels = [_wrap_label(c, 24) for c in display["categories"]]
        kind = spec["chart_type"]
        if kind == "bar":
            ax.bar(range(len(vals)), vals, color=[theme.palette[i % len(theme.palette)] for i in range(len(vals))], width=.64)
        elif kind == "line":
            ax.plot(range(len(vals)), vals, color=theme.accent, linewidth=3, marker="o", markersize=6)
        elif kind == "pie":
            valid = [(c, v) for c, v in zip(labels, spec["values"]) if v is not None and v > 0]
            pie_result = ax.pie([v for _, v in valid], colors=theme.palette,
                                startangle=90, autopct="%1.0f%%" if len(valid) <= 10 else None,
                                pctdistance=.75, wedgeprops={"width": .45, "edgecolor": theme.background},
                                textprops={"color": theme.ink, "fontsize": 10})
            wedges = pie_result[0]
            ax.legend(wedges, [c for c, _ in valid], loc="center left", bbox_to_anchor=(1, .5),
                      frameon=False, fontsize=max(5, 11 - len(valid) // 8), labelcolor=theme.ink,
                      ncols=1 if len(valid) <= 20 else 2)
            ax.set_aspect("equal")
        else:
            ax.scatter(spec["x_values"], vals, c=theme.accent, s=64, alpha=.85, edgecolors=theme.background)
        if kind in {"bar", "line"}:
            ax.set_xticks(range(len(labels)), labels, rotation=60 if len(labels) > 20 else 40 if len(labels) > 7 else 0,
                          ha="right" if len(labels) > 7 else "center", fontsize=max(6, 11 - len(labels) // 12))
        if kind != "pie":
            for side in ("top", "right", "left"):
                ax.spines[side].set_visible(False)
            ax.grid(axis="y", color=theme.grid, linewidth=.8)
            ax.set_axisbelow(True)
            ax.set_xlabel(display["x_label"], labelpad=12)
            ax.set_ylabel(display["y_label"], labelpad=12)
            ax.tick_params(axis="both", length=0, pad=8)
        missing = sum(v is None for v in spec["values"])
        if missing:
            fig.text(.95, .06, f"Missing observations: {missing}", ha="right", fontsize=9, color=theme.muted)
        files = _save_figure(fig, root, asset_id, display["title"])
    return {"id": asset_id, "kind": "chart", "title": display["title"], "files": files, "spec": spec}


def _diagram_spec(spec: Any) -> dict:
    if not isinstance(spec, dict):
        raise ValueError("Diagram specification must be an object")
    layout = spec.get("layout", "flow")
    if not isinstance(layout, str) or layout not in {"flow", "hierarchy", "timeline"}:
        raise ValueError("Unsupported diagram layout")
    nodes, ids = [], set()
    for node in _sequence(spec.get("nodes"), "nodes", 20, minimum=1):
        if not isinstance(node, dict):
            raise ValueError("Each node must be an object")
        node_id = _text(node.get("id"), "node id", 64, empty=False)
        if node_id in ids:
            raise ValueError("Diagram node IDs must be unique")
        ids.add(node_id)
        nodes.append({"id": node_id, "label": _text(node.get("label"), "node label", 160, empty=False)})
    edges = []
    for edge in _sequence(spec.get("edges", []), "edges", 30):
        if (not isinstance(edge, dict) or not isinstance(edge.get("source"), str)
                or not isinstance(edge.get("target"), str) or edge["source"] not in ids or edge["target"] not in ids):
            raise ValueError("Diagram edges must refer to existing nodes")
        edges.append({"source": edge["source"], "target": edge["target"],
                      "label": _text(edge.get("label", ""), "edge label", 80)})
    return {"title": _text(spec.get("title", "Diagram"), "title", 180, empty=False),
            "layout": layout, "nodes": nodes, "edges": edges}


def _positions(spec: dict) -> dict[str, tuple[float, float, float, float]]:
    nodes = spec["nodes"]
    if spec["layout"] == "hierarchy":
        parents = {node["id"]: set() for node in nodes}
        for edge in spec["edges"]:
            parents[edge["target"]].add(edge["source"])
        levels, remaining = [], set(parents)
        while remaining:
            layer = [n["id"] for n in nodes if n["id"] in remaining and not parents[n["id"]] & remaining]
            if not layer:
                raise ValueError("Hierarchy diagrams must not contain cycles")
            levels.append(layer)
            remaining.difference_update(layer)
        # Split exceptionally wide levels into rows so every box stays visible.
        rows = [layer[start:start + 5] for layer in levels for start in range(0, len(layer), 5)]
    else:
        columns = min(4 if spec["layout"] == "flow" else 5, len(nodes))
        rows = [[n["id"] for n in nodes[start:start + columns]] for start in range(0, len(nodes), columns)]
    positions = {}
    row_h = 1 / len(rows)
    for row_i, row in enumerate(rows):
        width = min(.25, .72 / len(row))
        height = min(.26, row_h * .6)
        for col_i, node_id in enumerate(row):
            positions[node_id] = ((col_i + .5) / len(row) - width / 2,
                                  1 - (row_i + .5) * row_h - height / 2, width, height)
    return positions


def _endpoints(a, b):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    ac, bc = (ax + aw / 2, ay + ah / 2), (bx + bw / 2, by + bh / 2)
    dx, dy = bc[0] - ac[0], bc[1] - ac[1]
    if abs(dx) > abs(dy):
        sign = 1 if dx >= 0 else -1
        return (ac[0] + sign * aw / 2, ac[1]), (bc[0] - sign * bw / 2, bc[1])
    sign = 1 if dy >= 0 else -1
    return (ac[0], ac[1] + sign * ah / 2), (bc[0], bc[1] - sign * bh / 2)


def render_diagram(spec, outdir, style="business", asset_id="diagram_1") -> dict:
    theme, spec, asset_id = get_style(style), _diagram_spec(spec), _identifier(asset_id)
    positions = _positions(spec)
    root = _output_dir(outdir)
    with matplotlib.rc_context(_rc(theme)):
        fig = Figure(figsize=(12, 6.75), facecolor=theme.background)
        FigureCanvasAgg(fig)
        ax = fig.add_axes([.055, .09, .89, .72])
        ax.set(xlim=(0, 1), ylim=(0, 1))
        ax.axis("off")
        fig.text(.065, .93, "\n".join(textwrap.wrap(spec["title"], 52)), va="top", fontsize=22, weight="bold")
        for edge in spec["edges"]:
            a, b = _endpoints(positions[edge["source"]], positions[edge["target"]])
            ax.add_patch(FancyArrowPatch(a, b, arrowstyle="-|>", mutation_scale=15, color=theme.secondary,
                                        linewidth=1.8, connectionstyle="arc3,rad=.08", zorder=1))
            if edge["label"]:
                ax.text((a[0] + b[0]) / 2, (a[1] + b[1]) / 2 + .025,
                        "\n".join(textwrap.wrap(edge["label"], 16)), ha="center", fontsize=8,
                        color=theme.muted, bbox={"facecolor": theme.background, "edgecolor": "none", "pad": 1})
        for index, node in enumerate(spec["nodes"]):
            x, y, width, height = positions[node["id"]]
            ax.add_patch(FancyBboxPatch((x, y), width, height, boxstyle="round,pad=.012,rounding_size=.016",
                                       facecolor=theme.surface, edgecolor=theme.accent, linewidth=1.6, zorder=2))
            text = "\n".join(textwrap.wrap(node["label"], max(8, int(width * 72))))
            ax.text(x + width / 2, y + height / 2, text, ha="center", va="center", zorder=3,
                    fontsize=max(5, min(12, 65 * height / max(1, text.count("\n") + 1))), color=theme.ink)
            if spec["layout"] == "timeline":
                ax.text(x + .012, y + height + .032, f"{index + 1:02d}", fontsize=10, color=theme.accent)
        files = _save_figure(fig, root, asset_id, spec["title"])
    return {"id": asset_id, "kind": "diagram", "title": spec["title"], "files": files, "spec": spec}


def _section_text(section: dict) -> str:
    body = "\n\n".join(section["paragraphs"])
    if section["bullets"]:
        body += ("\n\n" if body else "") + "\n".join("• " + b for b in section["bullets"])
    return body


def _fit_text(text: str, width: float, height: float, preferred: int, minimum: int = 6) -> int:
    """Conservative CJK-aware line budget; rejects content that cannot fit."""
    paragraphs = text.split("\n")
    for size in range(int(preferred), minimum - 1, -1):
        line_units = max(1, (width * 72 - 8) / (size * .58))
        lines = sum(max(1, math.ceil(sum(2 if ord(c) > 255 else 1 for c in p) / line_units))
                    for p in paragraphs)
        if lines * size * 1.2 + len(paragraphs) * size * .2 <= height * 72 - 5:
            return size
    raise ValueError("Text cannot fit this slide layout; shorten the section or divide it into sections")


def _blueprint(blueprint: Any) -> dict:
    if not isinstance(blueprint, dict):
        raise ValueError("Document blueprint must be an object")
    result = {"title": _text(blueprint.get("title"), "title", 180, empty=False),
              "subtitle": _text(blueprint.get("subtitle", ""), "subtitle", 300), "sections": []}
    for section in _sequence(blueprint.get("sections", []), "sections", 11):
        if not isinstance(section, dict):
            raise ValueError("Document sections must be objects")
        normalized = {"heading": _text(section.get("heading"), "heading", 160, empty=False),
                      "paragraphs": [_text(s, "paragraph", 600) for s in _sequence(section.get("paragraphs", []), "paragraphs", 4)],
                      "bullets": [_text(s, "bullet", 200) for s in _sequence(section.get("bullets", []), "bullets", 6)],
                      "asset_ids": [_identifier(s) for s in _sequence(section.get("asset_ids", []), "asset_ids", 2)]}
        total = sum(map(len, normalized["paragraphs"] + normalized["bullets"]))
        if total > (1600 if normalized["asset_ids"] else 2400):
            raise ValueError("Section text exceeds the bounded slide capacity")
        if len(set(normalized["asset_ids"])) != len(normalized["asset_ids"]):
            raise ValueError("A section cannot repeat an asset")
        _fit_text(_section_text(normalized), 4.25 if normalized["asset_ids"] else 11.7, 5.15, 21, minimum=11)
        result["sections"].append(normalized)
    return result


def _assets(assets: Any, root: Path) -> dict:
    if not isinstance(assets, dict) or len(assets) > 6:
        raise ValueError("A document supports at most six chart or diagram assets")
    clean = {}
    for asset_id, asset in assets.items():
        _identifier(asset_id)
        if not isinstance(asset, dict) or asset.get("id") != asset_id:
            raise ValueError("Asset IDs must match the asset map")
        if asset.get("kind") not in {"chart", "diagram"}:
            raise ValueError("Unsupported document asset")
        spec = _chart_spec(asset.get("spec")) if asset["kind"] == "chart" else _diagram_spec(asset.get("spec"))
        if asset["kind"] == "diagram":
            _positions(spec)
        png = None
        for item in _sequence(asset.get("files"), "asset files", 4, minimum=1):
            if not isinstance(item, dict) or not isinstance(item.get("path"), str):
                raise ValueError("Invalid asset file")
            path = Path(item["path"])
            if not path.is_absolute():
                path = root / path
            try:
                resolved = path.resolve(strict=True)
                resolved.relative_to(root)
            except (ValueError, OSError) as exc:
                raise ValueError("Asset files must stay inside this job's output directory") from exc
            if path.is_symlink() or not resolved.is_file() or resolved.stat().st_size > 12 * 1024 * 1024:
                raise ValueError("Unsafe or oversized asset file")
            if item.get("mime") == "image/png":
                if resolved.suffix.lower() != ".png":
                    raise ValueError("PNG assets require the .png extension")
                with Image.open(resolved) as image:
                    if image.format != "PNG" or image.width * image.height > 12_000_000:
                        raise ValueError("Invalid or oversized PNG asset")
                    image.verify()
                png = resolved
        if png is None:
            raise ValueError("A document asset must include a valid PNG")
        clean[asset_id] = {"id": asset_id, "kind": asset["kind"],
                           "title": spec["display"]["title"] if asset["kind"] == "chart" else spec["title"],
                           "spec": spec, "png": png}
    return clean


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.removeprefix("#"))


def _ppt_text(slide, text, x, y, w, h, theme, size=18, color=None, bold=False):
    size = _fit_text(text, w, h, size, minimum=min(6, int(size))) if text else size
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.02)
    frame.margin_top = frame.margin_bottom = Inches(.015)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = theme.font_family
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = _rgb(color or theme.ink)
        paragraph.space_after = Pt(size * .2)
        paragraph.line_spacing = 1.2
        # Explicit East Asian font avoids Office theme fallback for Korean.
        for run in paragraph.runs:
            rpr = run._r.get_or_add_rPr()
            ea = PptxElement("a:ea")
            ea.set("typeface", theme.font_family)
            rpr.append(ea)
    return box


def _ppt_base(prs, theme, number, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(theme.background)
    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(.16), prs.slide_height)
    accent.fill.solid()
    accent.fill.fore_color.rgb = _rgb(theme.accent)
    accent.line.fill.background()
    footer = notes[0] if notes else "AI Document Studio"
    if len(footer) > 120:
        footer = footer[:117] + "…"
    _ppt_text(slide, footer, .6, 7.02, 11.4, .3, theme, 8, theme.muted)
    _ppt_text(slide, str(number), 12.2, 7.0, .5, .3, theme, 10, theme.muted)
    if notes:
        slide.notes_slide.notes_text_frame.text = "Sources / 출처\n" + "\n".join(notes)
    return slide


class _LiteralWorkbookMixin:
    @property
    def xlsx_blob(self):
        """Keep untrusted labels literal in the native chart's Excel workbook.

        python-pptx's default writer enables XlsxWriter formula/URL inference.
        Its pinned writer still supplies range-compatible worksheet contents;
        only workbook options change here.
        """
        stream = io.BytesIO()
        with Workbook(stream, {"in_memory": True, "strings_to_formulas": False,
                               "strings_to_urls": False}) as workbook:
            sheet = workbook.add_worksheet()
            self._workbook_writer._populate_worksheet(workbook, sheet)
            sheet.write_string(0, 0, getattr(self, "_x_label", ""))
            originals = getattr(self, "_source_labels", ())
            if originals:
                source_sheet = workbook.add_worksheet("Source labels")
                source_sheet.write_row(0, 0, ["Field", "Item", "Part", "Original label"])
                source_sheet.set_column(0, 0, 16)
                source_sheet.set_column(3, 3, 80)
                row = 1
                for field, index, original in originals:
                    # Excel limits strings to 32,767 characters per cell. Split
                    # explicitly instead of accepting XlsxWriter's truncation.
                    for part, start in enumerate(range(0, max(1, len(original)), 32767), start=1):
                        source_sheet.write_string(row, 0, field)
                        source_sheet.write_number(row, 1, index)
                        source_sheet.write_number(row, 2, part)
                        source_sheet.write_string(row, 3, original[start:start + 32767])
                        row += 1
        return stream.getvalue()


class _LiteralCategoryData(_LiteralWorkbookMixin, CategoryChartData):
    pass


class _LiteralXyData(_LiteralWorkbookMixin, XyChartData):
    pass


def _ppt_chart(slide, asset, box, theme):
    spec, kind = asset["spec"], asset["spec"]["chart_type"]
    display = spec["display"]
    if kind == "scatter":
        data = _LiteralXyData()
        series = data.add_series(display["y_label"] or display["title"])
        for x, y in zip(spec["x_values"], spec["values"]):
            if y is not None:
                series.add_data_point(x, y)
        chart_type = XL_CHART_TYPE.XY_SCATTER
    else:
        data = _LiteralCategoryData()
        data.categories = display["categories"]
        data.add_series(display["y_label"] or display["title"], spec["values"])
        chart_type = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
                      "pie": XL_CHART_TYPE.DOUGHNUT}[kind]
        # Use standard pie XML for broad Office editor compatibility.
        if kind == "pie":
            chart_type = XL_CHART_TYPE.PIE
    data._x_label = display["x_label"]
    if display["shortened"]:
        data._source_labels = [("category", index + 1, label) for index, label in enumerate(spec["categories"])]
        data._source_labels += [(field, 0, spec[field]) for field in ("x_label", "y_label", "title")]
        data._source_labels.append(("source_caption", 0, _source_caption(spec)))
    x, y, w, h = box
    if display["shortened"]:
        _ppt_text(slide, "Labels shortened (…); see Source labels in chart data.", x, y + h - .24, w, .24,
                  theme, 7, theme.muted)
        h -= .28
    chart = slide.shapes.add_chart(chart_type, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    chart.font.name = theme.font_family
    chart.font.size = Pt(11)
    chart.font.color.rgb = _rgb(theme.ink)
    chart.has_title = False
    chart.has_legend = kind == "pie"
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.RIGHT
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(9 if len(spec["values"]) < 15 else 6)
        chart.legend.font.color.rgb = _rgb(theme.ink)
    chart.chart_style = 10
    if kind != "pie":
        for axis, label in ((chart.category_axis, display["x_label"]), (chart.value_axis, display["y_label"])):
            axis.tick_labels.font.name = theme.font_family
            axis.tick_labels.font.size = Pt(9)
            axis.tick_labels.font.color.rgb = _rgb(theme.muted)
            axis.has_title = bool(label)
            if label:
                axis.axis_title.text_frame.text = label
                for p in axis.axis_title.text_frame.paragraphs:
                    p.font.name = theme.font_family
                    p.font.size = Pt(10)
                    p.font.color.rgb = _rgb(theme.muted)
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = _rgb(theme.grid)
    for series in chart.series:
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = _rgb(theme.accent)
        series.format.line.color.rgb = _rgb(theme.accent)
        if kind in {"bar", "pie"}:
            for index, point in enumerate(series.points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = _rgb(theme.palette[index % len(theme.palette)])
                point.format.line.fill.background()
    # Native chart areas inherit the slide background instead of opaque white.
    for parent in (chart._chartSpace, chart._chartSpace.chart.plotArea):
        sppr = parent.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}spPr")
        if sppr is None:
            sppr = PptxElement("c:spPr")
            if parent is chart._chartSpace:
                parent.insert_element_before(sppr, "c:txPr", "c:externalData", "c:printSettings", "c:userShapes", "c:extLst")
            else:
                parent.insert_element_before(sppr, "c:extLst")
        fill = PptxElement("a:solidFill")
        color = PptxElement("a:srgbClr")
        color.set("val", theme.background[1:])
        fill.append(color)
        sppr.append(fill)
    return chart


def _ppt_diagram(slide, asset, box, theme):
    spec, positions = asset["spec"], _positions(asset["spec"])
    bx, by, bw, bh = box
    def transform(point):
        return Inches(bx + point[0] * bw), Inches(by + (1 - point[1]) * bh)
    for edge in spec["edges"]:
        start, end = _endpoints(positions[edge["source"]], positions[edge["target"]])
        conn = slide.shapes.add_connector(MSO_CONNECTOR_TYPE.STRAIGHT, *transform(start), *transform(end))
        conn.line.color.rgb = _rgb(theme.secondary)
        conn.line.width = Pt(1.6)
        arrow = PptxElement("a:tailEnd")
        arrow.set("type", "triangle")
        conn.line._get_or_add_ln().append(arrow)
        if edge["label"]:
            _ppt_text(slide, edge["label"], bx + (start[0] + end[0]) * bw / 2 - .5,
                      by + (1 - (start[1] + end[1]) / 2) * bh - .25, 1.0, .45, theme, 8, theme.muted)
    for node in spec["nodes"]:
        x, y, w, h = positions[node["id"]]
        px, py = bx + x * bw, by + (1 - y - h) * bh
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(px), Inches(py), Inches(w * bw), Inches(h * bh))
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(theme.surface)
        shape.line.color.rgb = _rgb(theme.accent)
        shape.line.width = Pt(1.3)
        frame = shape.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.margin_left = frame.margin_right = Inches(.07)
        frame.margin_top = frame.margin_bottom = Inches(.04)
        frame.text = node["label"]
        for p in frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            p.font.name = theme.font_family
            p.font.size = Pt(_fit_text(node["label"], w * bw, h * bh, 15, minimum=3))
            p.font.color.rgb = _rgb(theme.ink)


def _write_pptx(blueprint, assets, path, theme, notes):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = blueprint["title"]
    prs.core_properties.subject = "AI Document Studio"
    slide = _ppt_base(prs, theme, 1, notes)
    _ppt_text(slide, "DOCUMENT STUDIO", .7, 1.1, 11.8, .5, theme, 12, theme.accent, True)
    _ppt_text(slide, blueprint["title"], .7, 2.0, 11.7, 2.8, theme, theme.title_size + 7, bold=True)
    _ppt_text(slide, blueprint["subtitle"], .74, 5.15, 11.5, 1.2, theme, 20, theme.muted)
    for index, section in enumerate(blueprint["sections"], start=2):
        slide = _ppt_base(prs, theme, index, notes)
        _ppt_text(slide, section["heading"], .65, .5, 12, 1.0, theme, theme.title_size, bold=True)
        ids = section["asset_ids"]
        body = _section_text(section)
        if body:
            width = 4.25 if ids else 11.7
            # Bound point size by approximate line capacity, including CJK width.
            units = sum(2 if ord(c) > 255 else 1 for c in body)
            size = min(theme.body_size, max(11, int(math.sqrt(width * 5.0 * 3900 / max(units, 1)))))
            _ppt_text(slide, body, .7, 1.65, width, 5.15, theme, size)
        if ids:
            left, width = (5.3, 7.3) if body else (.75, 11.9)
            height = 4.8 / len(ids)
            for asset_index, asset_id in enumerate(ids):
                asset = assets[asset_id]
                top = 1.65 + asset_index * (height + .16)
                _ppt_text(slide, asset["title"], left, top, width, .45, theme, 14, theme.accent, True)
                box = (left, top + .48, width, height - .45)
                if asset["kind"] == "chart":
                    _ppt_chart(slide, asset, box, theme)
                else:
                    _ppt_diagram(slide, asset, box, theme)
        slide.notes_slide.notes_text_frame.text += "\n\n" + section["heading"] + "\n" + body
    prs.save(path)


def _word_font(style, theme, size, color):
    style.font.name = theme.font_family
    style.font.size = WordPt(size)
    style.font.color.rgb = WordRGBColor.from_string(color[1:])
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), theme.font_family)


def _write_docx(blueprint, assets, path, theme, notes):
    doc = Document()
    doc.core_properties.title = blueprint["title"]
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = WordInches(.72)
    section.left_margin = section.right_margin = WordInches(.8)
    # Word remains print-friendly; the selected theme supplies headings and rules.
    ink = "#182B49" if theme.id == "aws-desk" else theme.ink
    accent = "#087F83" if theme.id == "aws-desk" else theme.accent
    for name, size, color in (("Normal", 10.5, ink), ("Title", theme.title_size, ink),
                              ("Subtitle", 14, "#52647B"), ("Heading 1", 20, accent),
                              ("Heading 2", 14, accent), ("Caption", 9, "#52647B")):
        _word_font(doc.styles[name], theme, size, color)
    doc.styles["Normal"].paragraph_format.space_after = WordPt(8)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.25
    doc.add_heading(blueprint["title"], 0)
    if blueprint["subtitle"]:
        doc.add_paragraph(blueprint["subtitle"], style="Subtitle")
    for item in blueprint["sections"]:
        doc.add_heading(item["heading"], level=1)
        for paragraph in item["paragraphs"]:
            doc.add_paragraph(paragraph)
        for bullet in item["bullets"]:
            doc.add_paragraph(bullet, style="List Bullet")
        for asset_id in item["asset_ids"]:
            asset = assets[asset_id]
            doc.add_picture(str(asset["png"]), width=WordInches(6.6))
            doc.add_paragraph(asset["title"], style="Caption")
            if asset["kind"] == "chart" and asset["spec"]["display"]["shortened"]:
                doc.add_paragraph(_LABEL_NOTICE, style="Caption")
    if notes:
        doc.add_heading("Sources / 출처", level=1)
        for note in notes:
            doc.add_paragraph(note)
    footer = section.footer.paragraphs[0]
    footer.text = ((notes[0][:140] + ("…" if len(notes[0]) > 140 else "")) if notes else "AI Document Studio")
    footer.style = doc.styles["Caption"]
    doc.save(path)


def _write_preview(blueprint, assets, path, theme, notes):
    esc = html.escape
    content = [f'<header><p class="eyebrow">DOCUMENT STUDIO</p><h1>{esc(blueprint["title"])}</h1>',
               f'<p class="subtitle">{esc(blueprint["subtitle"])}</p></header>',
               '<aside class="notice">Content preview / 콘텐츠 미리보기. '
               'PowerPoint and Word may use different page breaks, fonts, and layouts. '
               '실제 Office 파일의 페이지 나눔과 배치는 다를 수 있습니다.</aside>']
    for section in blueprint["sections"]:
        content.append(f'<section><h2>{esc(section["heading"])}</h2>')
        content.extend(f'<p>{esc(text)}</p>' for text in section["paragraphs"])
        if section["bullets"]:
            content.append("<ul>" + "".join(f"<li>{esc(text)}</li>" for text in section["bullets"]) + "</ul>")
        for asset_id in section["asset_ids"]:
            asset = assets[asset_id]
            data = base64.b64encode(asset["png"].read_bytes()).decode("ascii")
            notice = " " + _LABEL_NOTICE if asset["kind"] == "chart" and asset["spec"]["display"]["shortened"] else ""
            content.append(f'<figure><img src="data:image/png;base64,{data}" alt="{esc(asset["title"], quote=True)}">'
                           f'<figcaption>{esc(asset["title"] + notice)}</figcaption></figure>')
        content.append("</section>")
    if notes:
        content.append('<footer><h2>Sources / 출처</h2>' + "".join(f"<p>{esc(note)}</p>" for note in notes) + "</footer>")
    font_css = ""
    candidates = [os.environ.get("STUDIO_FONT_PATH"), "/opt/studio/fonts/NanumSquareB.otf",
                  str(Path(__file__).resolve().parents[2] / "webapp/static/fonts/NanumSquareB.otf")]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            font_data = base64.b64encode(Path(candidate).read_bytes()).decode("ascii")
            font_css = (f'@font-face{{font-family:"{theme.font_family}";font-style:normal;'
                        f'font-weight:100 900;src:url(data:font/otf;base64,{font_data}) format("opentype")}}')
            break
    css = font_css + f'''*{{box-sizing:border-box}}body{{margin:0;background:{theme.background};color:{theme.ink};font:16px/1.7 "{theme.font_family}",sans-serif}}main{{max-width:1040px;margin:auto;padding:56px 32px}}header{{padding:24px 0 40px;border-bottom:3px solid {theme.accent}}}.eyebrow{{color:{theme.accent};font-size:12px;letter-spacing:.18em}}h1{{font-size:clamp(30px,5vw,{theme.title_size + 12}px);line-height:1.3;margin:14px 0;overflow-wrap:anywhere}}h2{{font-size:26px;line-height:1.4;color:{theme.accent};overflow-wrap:anywhere}}.subtitle,figcaption{{color:{theme.muted}}}.notice{{margin:24px 0;padding:16px 20px;background:{theme.surface};border-left:3px solid {theme.secondary};font-size:13px}}section{{padding:28px 0;border-bottom:1px solid {theme.grid}}}p,li{{white-space:pre-wrap;overflow-wrap:anywhere}}figure{{margin:28px 0}}img{{display:block;width:100%;height:auto;border-radius:10px}}figcaption{{font-size:13px;margin-top:8px}}footer{{padding-top:30px;font-size:13px}}@media(max-width:600px){{main{{padding:24px 18px}}}}'''
    document = '<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
    document += '<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; img-src data:; font-src data:; style-src \'unsafe-inline\'; base-uri \'none\'; form-action \'none\'">'
    document += f'<title>{esc(blueprint["title"])}</title><style>{css}</style></head><body><main>' + "".join(content) + "</main></body></html>"
    path.write_text(document, encoding="utf-8")


def render_documents(blueprint, assets, outdir, style="business", formats=("pptx", "docx"), source_notes=()) -> dict:
    theme, blueprint = get_style(style), _blueprint(blueprint)
    formats = _sequence(formats, "formats", 2)
    if any(not isinstance(f, str) or f not in {"pptx", "docx"} for f in formats) or len(set(formats)) != len(formats):
        raise ValueError("Document formats must be pptx and/or docx")
    notes = [_text(s, "source note", 600) for s in _sequence(source_notes, "source notes", 16)]
    root = _output_dir(outdir)
    assets = _assets(assets, root)
    used = {a for section in blueprint["sections"] for a in section["asset_ids"]}
    if not used <= set(assets):
        raise ValueError("Document references an unknown asset ID")
    destinations = {fmt: _target(root, f"documents.{fmt}") for fmt in formats}
    preview = _target(root, "preview.html")
    files = []
    if "pptx" in formats:
        _write_pptx(blueprint, assets, destinations["pptx"], theme, notes)
        files.append(_file(destinations["pptx"], PPTX_MIME))
    if "docx" in formats:
        _write_docx(blueprint, assets, destinations["docx"], theme, notes)
        files.append(_file(destinations["docx"], DOCX_MIME))
    _write_preview(blueprint, assets, preview, theme, notes)
    files.append(_file(preview, "text/html"))
    return {"id": "documents", "kind": "document", "title": blueprint["title"], "files": files, "spec": blueprint}
